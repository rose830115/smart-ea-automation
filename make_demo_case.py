#!/usr/bin/env python3
"""產生一組完全虛構的 Smart EA 測試案件（業務 Excel + 填好的範本 Excel）。

用途：從頭測整條流程，不必動到真實客戶案件。裡面沒有任何客戶資料。

案件設計成**四個倉間**（RMW / PL1 / PL2 / FGW），因為工廠有兩條不同產品的
產線；數據刻意做出高低差，評語才有東西可寫：

    PL1  最高風險：RH 72-78%、CO2 超標、PM10 高、多數物件含水超標
    RMW  次高：RH 偏高、部分紙箱與棧板含水超標
    PL2  中等：RH 貼近 60% 標準
    FGW  最低：全部合格

    .venv/bin/python make_demo_case.py            # 產到 demo_case/
    .venv/bin/python make_demo_case.py --outdir X # 指定位置
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

CASE_NAME = "DEMO 示範工廠"

# ── 區域配置：(area, function, zone, [checking points]) ──────────────────────
AREAS = [
    ("1", "Incoming Area", "RMW", ["1"]),
    ("2", "Storage Area", "RMW", ["2", "3"]),
    ("3", "Cutting Area", "PL1", ["4", "5"]),
    ("4", "Assembly", "PL1", ["6"]),
    ("5", "Printing Area", "PL2", ["7", "8"]),
    ("6", "Laminating Area", "PL2", ["9"]),
    ("7", "Dry Room", "FGW", ["10"]),
    ("8", "Storage Area", "FGW", ["11", "12"]),
]

# ── 環境數據：cp → (temp, humidity, co2, wind, pm10) ────────────────────────
ENV = {
    # RMW：濕度偏高
    "1":  (28.5, 66.2, 520, 0.35, 38.0),
    "2":  (28.1, 68.4, 545, 0.28, 41.5),
    "3":  (28.8, 65.1, 510, 0.42, 36.2),
    # PL1：最高風險（濕度高、CO2 超標、風速低、PM10 高）
    "4":  (31.2, 74.6, 685, 0.18, 62.4),
    "5":  (31.8, 77.3, 712, 0.12, 68.1),
    "6":  (30.9, 72.1, 648, 0.22, 58.7),
    # PL2：貼近標準
    "7":  (29.4, 61.5, 575, 0.46, 33.8),
    "8":  (29.1, 59.8, 552, 0.51, 31.2),
    "9":  (29.7, 62.3, 588, 0.44, 35.6),
    # FGW：合格
    "10": (26.2, 53.4, 470, 0.68, 22.1),
    "11": (26.8, 54.9, 485, 0.72, 24.3),
    "12": (26.5, 52.7, 462, 0.65, 21.8),
    # 室外
    "O1": (32.4, 79.5, 430, 0.85, 71.2),
    "O2": (32.1, 78.2, 425, 0.92, 68.9),
}

# ── 含水量：cp → [(物件, 用途, [量測值])] ───────────────────────────────────
MOISTURE = {
    "1":  [("Carton Box", "outer packaging", [12.8, 11.4, 13.1]),
           ("Wooden Pallet", "storage", [11.2, 10.7, 12.3])],
    "2":  [("Carton Box", "outer packaging", [13.5, 12.9, 14.2]),
           ("Leather", "raw material", [9.8, 9.2, 10.4]),
           ("Wooden Rack", "storage", [12.1, 11.8, 11.5])],
    "3":  [("Leather", "raw material", [8.9, 9.4, 8.7]),
           ("Fabric", "raw material", [9.1, 8.8, 9.6])],
    "4":  [("Semi-finished Goods", "in process", [13.2, 14.1, 12.8]),
           ("Table", "production equipment", [11.6, 12.2, 11.9]),
           ("Worker Tool Bag", "equipment", [10.8, 11.3, 10.5])],
    "5":  [("Semi-finished Goods", "in process", [14.6, 15.2, 13.9]),
           ("Wooden Pallet", "storage", [13.8, 14.3, 13.1]),
           ("Wall", "building structure", [12.4, 11.9, 12.8])],
    "6":  [("Semi-finished Goods", "in process", [12.9, 13.4, 12.2]),
           ("Shelf", "equipment", [11.1, 10.6, 11.8])],
    "7":  [("Ink", "raw material", [8.4, 8.9, 8.1]),
           ("Table", "production equipment", [9.6, 10.2, 9.8])],
    "8":  [("Semi-finished Goods", "in process", [10.4, 9.8, 10.9]),
           ("Carton Box", "outer packaging", [10.2, 9.6, 10.7])],
    "9":  [("Semi-finished Goods", "in process", [9.7, 10.3, 9.4]),
           ("Wall", "building structure", [9.9, 10.1, 9.5])],
    "10": [("Finished Goods", "finished product", [7.8, 8.2, 7.5]),
           ("Shelf", "equipment", [8.1, 7.9, 8.4])],
    "11": [("Finished Goods", "finished product", [8.4, 8.9, 8.2]),
           ("Carton Box", "outer packaging", [9.1, 8.7, 9.4])],
    "12": [("Finished Goods", "finished product", [7.9, 8.3, 8.0]),
           ("Wooden Pallet", "storage", [9.2, 8.8, 9.5])],
}

# ── 棉棒：cp → [(棉棒編號, 物件, 用途)] ─────────────────────────────────────
SWABS = {
    "1":  [("1", "Carton Box", "outer packaging"), ("2", "Leather", "raw material")],
    "2":  [("3", "Wooden Pallet", "storage"), ("4", "Leather", "raw material")],
    "3":  [("5", "Wooden Rack", "storage")],
    "4":  [("6", "Table", "production equipment"), ("7", "Leather", "raw material")],
    "5":  [("8", "Semi-finished Goods", "in process"), ("9", "Wooden Pallet", "storage"),
           ("10", "Wall", "building structure")],
    "6":  [("11", "Shelf", "equipment"), ("12", "Semi-finished Goods", "in process")],
    "7":  [("13", "Table", "production equipment")],
    "8":  [("14", "Semi-finished Goods", "in process"), ("15", "Worker Tool Bag", "equipment")],
    "9":  [("16", "Wall", "building structure")],
    "10": [("17", "Finished Goods", "finished product")],
    "11": [("18", "Shelf", "equipment"), ("19", "Carton Box", "outer packaging")],
    "12": [("20", "Finished Goods", "finished product")],
}

# ── 菌落計數：棉棒編號 → {菌種: 菌落數}。PL1 的棉棒刻意高 ──────────────────
COLONIES = {
    "1":  {"Aspergillus spp.": 3, "Penicillium spp.": 2},
    "2":  {"Cladosporium spp.": 1},
    "3":  {"Aspergillus spp.": 5, "Cladosporium spp.": 2},
    "4":  {"Penicillium spp.": 2},
    "5":  {"Aspergillus spp.": 4},
    "6":  {"Aspergillus spp.": 12, "Penicillium spp.": 8},
    "7":  {"Aspergillus spp.": 9, "Cladosporium spp.": 4},
    "8":  {"Penicillium spp.": 45, "Aspergillus spp.": 22},
    "9":  {"Aspergillus spp.": 38, "Cladosporium spp.": 15},
    "10": {"Penicillium spp.": 18},
    "11": {"Aspergillus spp.": 11, "Penicillium spp.": 6},
    "12": {"Aspergillus spp.": 14},
    "13": {"Cladosporium spp.": 3},
    "14": {"Penicillium spp.": 5, "Aspergillus spp.": 2},
    "15": {"Aspergillus spp.": 7},
    "16": {"Cladosporium spp.": 2},
    "17": {"Penicillium spp.": 1},
    "18": {"Aspergillus spp.": 2},
    "19": {"Cladosporium spp.": 1},
    "20": {},
}

# ── ITEMLIST：物件 → (Moisture 分類, CFU 分類)。分類碼與正式範本一致 ────────
ITEMLIST = [
    ("Leather", "a.原材料", "a.原材料"),
    ("Fabric", "a.原材料", "a.原材料"),
    ("Ink", "a.原材料", "a.原材料"),
    ("Semi-finished Goods", "b.半成品", "b.半成品"),
    ("Finished Goods", "c.成品", "c.成品"),
    ("Carton Box", "e.外包裝", "e.外包裝"),
    ("Wooden Pallet", "f.設備", "f.設備"),
    ("Wooden Rack", "f.設備", "f.設備"),
    ("Shelf", "f.設備", "f.設備"),
    ("Worker Tool Bag", "f.設備", "f.設備"),
    ("Table", "j.生產設備", "j.生產設備"),
    ("Wall", "g.建築本體", "g.建築本體"),
]

ZONE_SHEETS = [
    ("Raw Material warehouse", "RMW"),
    ("Production Line", "PL"),
    ("Finished Goods Warehouse", "FGW"),
]

VENDOR_HEADERS = [
    "Checking points", "Function", "Temp. (°C)", "Humidity (%)", "CO2 (ppm)",
    "Wind flow (m/s)", "PM10 (μg/m³)", "Moisture checking Item",
    "Specific function/material",
    "Moisture content (%) - 1", "Moisture content (%) - 2", "Moisture content (%) - 3",
    "Moisture content (%) - 4", "Moisture content (%) - 5",
    "Swab No.", "棉棒編號", "Swab Item", "Specific function/material",
]

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")


def zone_of_cp(cp: str) -> str:
    for _area, _fn, zone, cps in AREAS:
        if cp in cps:
            return zone
    return ""


def area_of_cp(cp: str) -> tuple[str, str]:
    for area, fn, _zone, cps in AREAS:
        if cp in cps:
            return area, fn
    return "", ""


def write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def build_vendor(path: Path) -> None:
    """業務回傳的環境資料 Excel（ver_5 版面）。"""
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, zone_base in ZONE_SHEETS:
        ws = wb.create_sheet(sheet_name)
        ws.append([sheet_name, "Every checking point should be filled in by the audit team."])
        write_header(ws, VENDOR_HEADERS)
        # 業務檔仍是三個分頁：兩條產線都在 Production Line 分頁，
        # 拆成 PL1 / PL2 是在範本的 Areas 分頁決定的。
        for cp in [c for c in ENV if not c.startswith("O")]:
            base = zone_of_cp(cp)
            if not base.startswith(zone_base):
                continue
            _area, function = area_of_cp(cp)
            temp, hum, co2, wind, pm10 = ENV[cp]
            moist = MOISTURE.get(cp, [])
            swabs = SWABS.get(cp, [])
            for i in range(max(len(moist), len(swabs), 1)):
                row = [None] * len(VENDOR_HEADERS)
                if i == 0:
                    row[0], row[1] = cp, function
                    row[2], row[3], row[4], row[5], row[6] = temp, hum, co2, wind, pm10
                if i < len(moist):
                    obj, use, values = moist[i]
                    row[7], row[8] = obj, use
                    for j, v in enumerate(values[:5]):
                        row[9 + j] = v
                if i < len(swabs):
                    swab_no, obj, use = swabs[i]
                    row[14] = f"{cp}.{i + 1}"
                    row[15] = int(swab_no)
                    row[16], row[17] = obj, use
                ws.append(row)

    ws = wb.create_sheet("Outside")
    ws.append(["Outside"])
    write_header(ws, ["Checking points", "Temp. (°C)", "Humidity (%)", "CO2 (ppm)",
                      "Wind flow (m/s)", "PM10 (μg/m³)"])
    for cp in [c for c in ENV if c.startswith("O")]:
        temp, hum, co2, wind, pm10 = ENV[cp]
        ws.append([cp, temp, hum, co2, wind, pm10])

    wb.save(path)


def build_target(path: Path, split: bool) -> None:
    """範本 / 參考檔：Areas 已填好，split=True 時 PL 拆成 PL1 / PL2。"""
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Areas")
    write_header(ws, ["Checking Point", "Report Area", None, "Report Area",
                      "Function", "RMW/PL/FGW/OUTSIDE", "Moisture Check", "CFU Check"])
    rows = []
    for area, function, zone, cps in AREAS:
        if not split and zone.startswith("PL"):
            zone = "PL"
        for cp in cps:
            rows.append([cp, area, None, None, None, None, None, None])
        rows[len(rows) - len(cps)][3] = area
        rows[len(rows) - len(cps)][4] = function
        rows[len(rows) - len(cps)][5] = zone
        rows[len(rows) - len(cps)][6] = True
        rows[len(rows) - len(cps)][7] = True
    for row in rows:
        ws.append(row)

    ws = wb.create_sheet("Env. Data")
    write_header(ws, ["Checking Point", "Report Area", "Main Zone", "Temp. (°C)",
                      "Humidity (%)", "CO2 (ppm)", "Wind Flow (m/s)", "PM10 (μg/m³)"])

    ws = wb.create_sheet("Moisture")
    write_header(ws, ["Checking Point", "Report Area", "Main Zone", "Classification",
                      "Object"] + [f"Moisture (%) - {i}" for i in range(1, 6)])

    ws = wb.create_sheet("CFU")
    write_header(ws, ["Checking Point", "Swab", "Report Area", "Main Zone",
                      "Classification", "Object", "Count", "CFU/m²"])

    ws = wb.create_sheet("菌落計數表")
    species = sorted({s for counts in COLONIES.values() for s in counts})
    swab_ids = sorted(COLONIES, key=int)
    ws.append(["示範案件菌落計數表"])
    ws.append(["菌種 \\ 棉棒編號"] + [int(s) for s in swab_ids])
    for sp in species:
        ws.append([sp] + [COLONIES[s].get(sp, 0) for s in swab_ids])
    ws.append(["Total CFU"] + [sum(COLONIES[s].values()) for s in swab_ids])

    ws = wb.create_sheet("ITEMLIST")
    write_header(ws, ["Object", "Moisture 分類", "CFU 分類"])
    for item in ITEMLIST:
        ws.append(list(item))

    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--outdir", type=Path,
                        default=Path(__file__).resolve().parent / "demo_case")
    args = parser.parse_args()
    args.outdir.mkdir(parents=True, exist_ok=True)

    vendor = args.outdir / "DEMO_業務回傳環境資料.xlsx"
    target4 = args.outdir / "DEMO_範本_四倉間(PL1+PL2).xlsx"
    target3 = args.outdir / "DEMO_範本_三倉間(對照組).xlsx"

    build_vendor(vendor)
    build_target(target4, split=True)
    build_target(target3, split=False)

    print(f"業務檔　　　　：{vendor}")
    print(f"範本（四倉間）：{target4}")
    print(f"範本（三倉間）：{target3}")
    print()
    print(f"室內 checking point {len([c for c in ENV if not c.startswith('O')])} 個、"
          f"室外 {len([c for c in ENV if c.startswith('O')])} 個、"
          f"area {len(AREAS)} 個、棉棒 {len(COLONIES)} 支")


if __name__ == "__main__":
    main()
