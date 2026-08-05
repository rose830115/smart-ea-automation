#!/usr/bin/env python3
"""既有案件回歸測試。

改動 smart_ea_automation.py 之前先跑一次存基準，改完再跑一次比對，
確認「這次修的 bug」沒有順手弄壞其他案件。

    .venv/bin/python regression.py --update    # 存基準（改 code 前）
    .venv/bin/python regression.py             # 比對（改 code 後）

案件與輸入檔從 uploads/ 自動探測，取每個案件最後一次上傳的那組。
基準存在 regression_baseline/，內含客戶資料，不進版控。
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
import warnings
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from smart_ea_automation import run_automation, run_from_standardized

warnings.filterwarnings("ignore", message="Data Validation extension is not supported")

PROJECT_DIR = Path(__file__).resolve().parent
UPLOAD_ROOT = PROJECT_DIR / "uploads"
BASELINE_DIR = PROJECT_DIR / "regression_baseline"

VENDOR_SHEETS = {"Raw Material warehouse", "Production Line", "Finished Goods Warehouse"}
TARGET_SHEETS = {"Areas", "ITEMLIST"}
STANDARDIZED_SHEETS = {"Env. Data", "Moisture", "CFU"}

# summary 這幾個數字是最快看出「整批資料掉了」的指標
SUMMARY_KEYS = [
    "env_rows",
    "moisture_rows",
    "microbiology_rows",
    "indoor_area_count",
    "quality_issue_count",
]


def sheets_of(path: Path) -> set[str]:
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            return set(wb.sheetnames)
        finally:
            wb.close()
    except Exception:
        return set()


def discover_cases() -> list[dict[str, Any]]:
    """從 uploads/ 找出每個案件最後一次上傳的輸入檔。"""
    latest: dict[str, tuple[int, Path]] = {}
    for d in UPLOAD_ROOT.iterdir() if UPLOAD_ROOT.exists() else []:
        if not d.is_dir():
            continue
        m = re.match(r"^(.*)_(\d{10})$", d.name)
        if not m or not list(d.glob("*.xlsx")):
            continue
        case, ts = m.group(1), int(m.group(2))
        if case not in latest or ts > latest[case][0]:
            latest[case] = (ts, d)

    cases = []
    for case, (_, d) in sorted(latest.items()):
        vendor = target = adjusted = None
        for path in sorted(d.glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            names = sheets_of(path)
            if VENDOR_SHEETS.issubset(names):
                vendor = path
            elif TARGET_SHEETS.issubset(names):
                target = path
            elif STANDARDIZED_SHEETS.issubset(names):
                adjusted = path

        if vendor:
            cases.append({"name": case, "mode": "fresh", "vendor": vendor, "target": target})
        elif adjusted:
            cases.append({"name": case, "mode": "standardized", "adjusted": adjusted})
    return cases


def fingerprint(payload: dict[str, Any]) -> dict[str, Any]:
    """把 payload 壓成可比對的結構。

    直接 diff 整包 JSON 雜訊太多（欄位順序、浮點尾數），改成只留
    「送進後台會影響結果」的東西：各 area 的量測值與分類。
    """
    fp: dict[str, Any] = {
        "summary": {k: payload.get("summary", {}).get(k) for k in SUMMARY_KEYS},
        "quality_issues": sorted(json.dumps(q, sort_keys=True, ensure_ascii=False)
                                 for q in payload.get("quality_issues", [])),
        "areas": {},
    }

    for zone in payload.get("indoor", []) or []:
        for area in zone.get("areas", []) or []:
            key = f"{zone.get('main_zone')}/{area.get('assessment_area')}"
            if key in fp["areas"]:  # key 撞了就是分區邏輯出問題，別默默蓋掉
                raise ValueError(f"payload 有重複的 area key：{key}")
            env = area.get("env", {}) or {}
            fp["areas"][key] = {
                "function": area.get("function"),
                "checking_points": sorted(map(str, area.get("checking_points") or [])),
                "env": {k: sorted(map(str, v)) for k, v in sorted(env.items()) if isinstance(v, list)},
                "moisture": sorted(
                    f"{m.get('checking_point')}|{m.get('category_to_select')}|{m.get('object')}|"
                    f"{sorted(map(str, m.get('values') or []))}"
                    for m in (area.get("moisture") or [])
                ),
                "microbiology": sorted(
                    f"{m.get('checking_point')}|{m.get('swab')}|{m.get('category_to_select')}|"
                    f"{m.get('object')}|{m.get('count_to_enter')}|"
                    f"{sorted(m.get('isolate_mold_species_names') or [])}"
                    for m in (area.get("microbiology") or [])
                ),
            }

    outside = payload.get("outside")
    if outside is not None:
        fp["outside"] = json.dumps(outside, sort_keys=True, ensure_ascii=False)

    fp["yims"] = yims_fingerprint(payload)
    return fp


def yims_fingerprint(payload: dict[str, Any]) -> dict[str, Any]:
    """比對真正送進後台的那一層。

    backend payload 對了不代表後台收到的對——含水率合併與分類排序都發生在
    這一層。卡片順序本身就是要驗的東西，所以這裡的 list 一律不排序。
    """
    from yims_payload_builder import build_yims_extends_data

    out: dict[str, Any] = {}
    extends = build_yims_extends_data(payload)
    for main_area in extends.get("main_area_array", []) or []:
        zone = json.dumps(main_area.get("main_area_name"), sort_keys=True, ensure_ascii=False)
        for area in main_area.get("assessment_area_array", []) or []:
            name = json.dumps(area.get("assessment_area_name"), sort_keys=True, ensure_ascii=False)
            out[f"{zone}/{name}"] = {
                "moisture_cards": [
                    {
                        "type": r.get("object_moisture_content_type"),
                        "values": r.get("object_moisture_contents"),
                        "avg": r.get("object_moisture_content_average"),
                        "fail_rate": r.get("object_moisture_content_failure_rate"),
                    }
                    for r in area.get("object_moisture_content_array", []) or []
                ],
                "moisture_total_fail_rate": area.get("object_moisture_content_total_failure_rate"),
                "microbiology_cards": [
                    {
                        "type": r.get("microbiology_sampling_object_type"),
                        "samples": [
                            [
                                o.get("microbiology_sampling_object_sampling_number"),
                                o.get("microbiology_sampling_object_sampling_value"),
                                o.get("microbiology_sampling_object_sampling_cfu"),
                            ]
                            for o in r.get("microbiology_sampling_objects", []) or []
                        ],
                        "cfu_avg": r.get("microbiology_sampling_object_sampling_cfu_average"),
                        "species": r.get("microbiology_sampling_object_isolate_mold_species"),
                    }
                    for r in area.get("microbiology_sampling_object_array", []) or []
                ],
            }
    return out


def diff(old: Any, new: Any, path: str = "") -> list[str]:
    if isinstance(old, dict) and isinstance(new, dict):
        out = []
        for k in sorted(set(old) | set(new)):
            sub = f"{path}.{k}" if path else str(k)
            if k not in old:
                out.append(f"  + 新增 {sub}")
            elif k not in new:
                out.append(f"  - 消失 {sub}")
            else:
                out.extend(diff(old[k], new[k], sub))
        return out
    if isinstance(old, list) and isinstance(new, list) and old != new:
        gone = [x for x in old if x not in new]
        added = [x for x in new if x not in old]
        out = []
        for x in gone[:3]:
            out.append(f"  - {path} 少了：{x}")
        for x in added[:3]:
            out.append(f"  + {path} 多了：{x}")
        if len(gone) > 3 or len(added) > 3:
            out.append(f"    …{path} 共少 {len(gone)} 筆、多 {len(added)} 筆")
        return out
    if old != new:
        return [f"  ~ {path}：{old!r} → {new!r}"]
    return []


def run_case(case: dict[str, Any], outdir: Path) -> dict[str, Any]:
    if case["mode"] == "fresh":
        return run_automation(
            vendor_path=case["vendor"],
            target_path=case.get("target"),
            outdir=outdir,
            case_name=case["name"],
        )
    return run_from_standardized(case["adjusted"], outdir, case["name"])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--update", action="store_true", help="把目前行為存成基準（改 code 之前跑）")
    ap.add_argument("--case", help="只跑指定案件")
    args = ap.parse_args()

    cases = discover_cases()
    if args.case:
        cases = [c for c in cases if c["name"] == args.case]
    if not cases:
        print("找不到可跑的案件（uploads/ 是空的？）")
        return 1

    BASELINE_DIR.mkdir(exist_ok=True)
    failures = 0
    tmp_root = Path(tempfile.mkdtemp(prefix="sea_regression_"))

    try:
        for case in cases:
            name = case["name"]
            label = "從業務 Excel" if case["mode"] == "fresh" else "從調整後 Excel"
            print(f"\n=== {name}（{label}）")

            try:
                result = run_case(case, tmp_root / name)
            except Exception as exc:
                print(f"  ✗ 執行就失敗了：{type(exc).__name__}: {exc}")
                failures += 1
                continue

            payload = json.loads(Path(result["paths"]["backend_payload"]).read_text(encoding="utf-8"))
            fp = fingerprint(payload)
            summary = " ".join(f"{k}={fp['summary'][k]}" for k in SUMMARY_KEYS)
            print(f"  {summary}")

            baseline_path = BASELINE_DIR / f"{name}.json"
            if args.update:
                baseline_path.write_text(
                    json.dumps(fp, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8"
                )
                print(f"  ✓ 已存基準 {baseline_path.name}")
                continue

            if not baseline_path.exists():
                print(f"  ⚠ 沒有基準可比（先跑一次 --update）")
                continue

            old = json.loads(baseline_path.read_text(encoding="utf-8"))
            changes = diff(old, fp)
            if changes:
                failures += 1
                print(f"  ✗ 與基準有 {len(changes)} 處差異：")
                for line in changes[:25]:
                    print(line)
                if len(changes) > 25:
                    print(f"    …另有 {len(changes) - 25} 處")
            else:
                print("  ✓ 與基準一致")
    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)

    print()
    if args.update:
        print(f"基準已更新（{len(cases)} 個案件）→ {BASELINE_DIR}")
        return 0
    if failures:
        print(f"✗ {failures} 個案件與基準不符——確認這些差異是這次修改「預期造成的」再繼續。")
        return 1
    print(f"✓ {len(cases)} 個案件全部與基準一致。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
