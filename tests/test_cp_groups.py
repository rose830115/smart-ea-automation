"""群組採樣點解析：拆分、同區通過、跨區衝突、成員查無。 .venv/bin/python -m unittest discover -s tests"""
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
warnings.filterwarnings("ignore")

from cp_groups import split_cp_group  # noqa: E402
from smart_ea_automation import ReferenceData, build_backend_payload, read_vendor, report_area_for_cp  # noqa: E402
from validation import validate_inputs  # noqa: E402


def make_ref(area_by_cp):
    info = {"Area 1": {"zone": "RMW", "function": "Storage"}, "Area 2": {"zone": "PL", "function": "Prod"}}
    return ReferenceData(area_by_cp, info, {}, {}, {}, {}, {})


AREAS = {"1": "Area 1", "2": "Area 1", "3": "Area 1", "10": "Area 2", "12": "Area 2"}


class SplitTest(unittest.TestCase):
    def test_separators(self):
        self.assertEqual(split_cp_group("1,2,3"), ["1", "2", "3"])
        self.assertEqual(split_cp_group("1，2，3"), ["1", "2", "3"])
        self.assertEqual(split_cp_group("1、2 、 3"), ["1", "2", "3"])
        self.assertEqual(split_cp_group("3b, 4"), ["3B", "4"])
        self.assertEqual(split_cp_group("7"), ["7"])
        self.assertEqual(split_cp_group("1,2,"), ["1", "2"])


class ResolveTest(unittest.TestCase):
    def test_same_area(self):
        ref = make_ref(AREAS)
        self.assertEqual(report_area_for_cp("3,1,2", ref, "RMW"), "Area 1")
        self.assertEqual(report_area_for_cp("1，2", ref, "RMW"), "Area 1")

    def test_cross_area_conflict_not_guessed(self):
        self.assertEqual(report_area_for_cp("1,10", make_ref(AREAS), "RMW"), "")

    def test_unknown_member_not_guessed(self):
        self.assertEqual(report_area_for_cp("1,99", make_ref(AREAS), "RMW"), "")

    def test_single_cp_unchanged(self):
        ref = make_ref(AREAS)
        self.assertEqual(report_area_for_cp("10", ref, "PL"), "Area 2")
        self.assertEqual(report_area_for_cp("99", ref, "PL"), "")


def build_vendor(path, cps):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Raw Material warehouse")
    ws.append(["RMW"])
    ws.append(["Checking points", "Function", "Temp. (°C)", "Humidity (%)", "CO2 (ppm)", "Wind flow (m/s)", "PM10 (μg/m³)"])
    for cp in cps:
        ws.append([cp, "Storage", 26.0, 70.0, 400, 0.3, 20])
    for name in ("Production Line", "Finished Goods Warehouse", "Outside"):
        w = wb.create_sheet(name)
        w.append([name])
        w.append(["Checking points", "Function", "Temp. (°C)"])
    wb.save(path)


class EndToEndTest(unittest.TestCase):
    def run_payload(self, cps):
        ref = make_ref(AREAS)
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "v.xlsx"
            build_vendor(p, cps)
            data = read_vendor(p, ref)
        return data, build_backend_payload(data, ref, "t")

    def test_group_rows_land_in_area_and_label_preserved(self):
        data, payload = self.run_payload(["1,2,3", None, "10,12"])
        self.assertEqual([r["Checking Point"] for r in data["Env. Data"]], ["1,2,3", "1,2,3", "10,12"])
        self.assertEqual([r["Report Area"] for r in data["Env. Data"]], ["Area 1", "Area 1", "Area 2"])
        self.assertEqual(payload["quality_issues"], [])
        self.assertEqual(payload["summary"]["indoor_area_count"], 2)

    def test_conflict_and_unknown_flagged_with_names(self):
        _, payload = self.run_payload(["1,10", "2,99"])
        kinds = {i["issue"]: i["detail"] for i in payload["quality_issues"] if i["issue"].startswith("group_")}
        self.assertIn("1→Area 1", kinds["group_area_conflict"])
        self.assertIn("10→Area 2", kinds["group_area_conflict"])
        self.assertIn("99", kinds["group_unknown_member"])
        self.assertEqual(payload["summary"]["indoor_area_count"], 0)

    def test_preflight_accepts_group_and_rejects_junk(self):
        with tempfile.TemporaryDirectory() as d:
            ok, bad = Path(d) / "ok.xlsx", Path(d) / "bad.xlsx"
            build_vendor(ok, ["1,2,3"])
            build_vendor(bad, ["Cutting area"])
            self.assertEqual([v for v in validate_inputs(ok, None) if "cp 欄位" in v], [])
            self.assertTrue([v for v in validate_inputs(bad, None) if "cp 欄位" in v])


if __name__ == "__main__":
    unittest.main()
